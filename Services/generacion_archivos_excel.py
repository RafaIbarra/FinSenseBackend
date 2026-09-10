from datetime import date
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter
from sqlalchemy.ext.asyncio import AsyncSession
from Schemas.file_format_schemas import ExcelIvaFormat
from Schemas.Respuestas import RespuestaFuncion
async def generar_excel_mes(db: AsyncSession,mes: int,anno:int ,data:ExcelIvaFormat):
    try:
        wb = Workbook()
        ws = wb.active
        ws.title = f"IVA {mes:02d}-{anno}"


        # directorio_script = Path(__file__).resolve().parent
        # nombre_archivo = f"IVA_{mes:02d}_{anno}.xlsx"
        # ruta_fisica = directorio_script / nombre_archivo
        # Encabezados obtenidos directamente del modelo Pydantic
        encabezados = list(ExcelIvaFormat.model_fields.keys())

        ws.append(encabezados)
        for cell in ws[1]:
                cell.font = Font(bold=True)
        for movimiento in data:
                datos = movimiento.model_dump()

                ws.append([
                    datos[campo]
                    for campo in encabezados
                ])
        # Formato de columnas
        columna_total_gasto = encabezados.index("TotalGasto") + 1
        columna_iva_diez = encabezados.index("IvaDiez") + 1
        columna_iva_cinco = encabezados.index("IvaCinco") + 1

        columna_fecha_factura = encabezados.index("FechaFactura") + 1
        columna_fecha_registro = encabezados.index("FechaRegistro") + 1

        # Formato numérico
        for fila in range(2, ws.max_row + 1):
            ws.cell(fila, columna_total_gasto).number_format = '#,##0'
            ws.cell(fila, columna_iva_diez).number_format = '#,##0'
            ws.cell(fila, columna_iva_cinco).number_format = '#,##0'

            ws.cell(fila, columna_fecha_factura).number_format = 'dd/mm/yyyy'
            ws.cell(fila, columna_fecha_registro).number_format = 'dd/mm/yyyy hh:mm:ss'

        # Totales
        fila_total = ws.max_row + 1

        ws.cell(fila_total, 1, "TOTAL")

        ultima_fila_datos = fila_total - 1
        ws.cell(
            fila_total,
            columna_total_gasto,
            f"=SUM({get_column_letter(columna_total_gasto)}2:{get_column_letter(columna_total_gasto)}{ultima_fila_datos})"
        )

        ws.cell(
            fila_total,
            columna_iva_diez,
            f"=SUM({get_column_letter(columna_iva_diez)}2:{get_column_letter(columna_iva_diez)}{ultima_fila_datos})"
        )

        ws.cell(
            fila_total,
            columna_iva_cinco,
            f"=SUM({get_column_letter(columna_iva_cinco)}2:{get_column_letter(columna_iva_cinco)}{ultima_fila_datos})"
        )

        # Formato de la fila de totales
        for cell in ws[fila_total]:
            cell.font = Font(bold=True)

        ws.cell(fila_total, columna_total_gasto).number_format = '#,##0'
        ws.cell(fila_total, columna_iva_diez).number_format = '#,##0'
        ws.cell(fila_total, columna_iva_cinco).number_format = '#,##0'

        # Autofiltro solamente sobre los datos
        ws.auto_filter.ref = f"A1:{get_column_letter(len(encabezados))}{ultima_fila_datos}"

        # Congelar encabezado
        ws.freeze_panes = "A2"

        # Ajustar ancho de columnas
        for columna in range(1, len(encabezados) + 1):
            letra = get_column_letter(columna)

            longitud = len(str(ws.cell(1, columna).value))

            for fila in range(2, ws.max_row + 1):
                valor = ws.cell(fila, columna).value

                if valor is not None:
                    longitud = max(longitud, len(str(valor)))

            ws.column_dimensions[letra].width = min(longitud + 2, 40)

        # Generar archivo en memoria
        # print(ruta_fisica)
        # wb.save(ruta_fisica)
        archivo = BytesIO()
        wb.save(archivo)
        archivo.seek(0)

        # return archivo
        return RespuestaFuncion(data_registro=archivo) 
    except Exception as e:
        return RespuestaFuncion(success_registro=False,mensaje=f"Error al generar el archivo Excel de IVA: {str(e)}")
        